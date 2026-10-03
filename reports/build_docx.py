"""Genera reports/Informe_Practica1_Wine_Reviews.docx a partir de los mismos bloques (story) que el PDF.

Sección 1 (A4, una columna): documento integrado. Sección 2-3 (Carta): artículo IEEE short paper
(título a una columna + cuerpo a dos columnas).
"""
import html
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor
from reportlab.platypus import Paragraph, Table, Image, KeepTogether, PageBreak, FrameBreak, NextPageTemplate, Spacer

sys.path.insert(0, str(Path(__file__).parent))
import build_report as br  # noqa: E402  (construye `story`; el PDF se genera sobre una copia)

FONTS = {"Ar": "Arial", "Ti": "Times New Roman", "Courier": "Courier New"}
ALIGN = {0: WD_ALIGN_PARAGRAPH.LEFT, 1: WD_ALIGN_PARAGRAPH.CENTER, 2: WD_ALIGN_PARAGRAPH.RIGHT, 4: WD_ALIGN_PARAGRAPH.JUSTIFY}
TAG = re.compile(r"(</?b>|</?i>|<br/>)")

doc = Document()
doc.core_properties.title = "Práctica 1 – Minería de Datos: Wine Reviews"
doc.core_properties.author = "Edison Paul Llerena Cuzco"
doc.core_properties.subject = "Predicción de la calidad de vinos a partir de reseñas de catadores"

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")


# ------------------------------------------------------------------ utilidades
def font_of(style):
    name = style.fontName
    fam, _, flag = name.partition("-")
    return FONTS.get(fam, "Arial"), "B" in flag, "I" in flag


def set_shading(element_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element_pr.append(shd)


def add_runs(par, markup, rstyle):
    """Convierte el mini-markup de reportlab (<b>, <i>, <br/>) en runs de Word."""
    fam, base_b, base_i = font_of(rstyle)
    bold, ital = base_b, base_i
    color = rstyle.textColor
    for tok in TAG.split(markup):
        if not tok:
            continue
        if tok == "<b>": bold = True; continue
        if tok == "</b>": bold = base_b; continue
        if tok == "<i>": ital = True; continue
        if tok == "</i>": ital = base_i; continue
        if tok == "<br/>":
            par.add_run().add_break(WD_BREAK.LINE); continue
        r = par.add_run(html.unescape(tok))
        r.font.name = fam
        r._element.rPr.rFonts.set(qn("w:eastAsia"), fam)
        r.font.size = Pt(rstyle.fontSize)
        r.bold, r.italic = bold, ital
        hv = color.hexval()[2:] if hasattr(color, "hexval") else "000000"
        r.font.color.rgb = RGBColor.from_string(hv.upper())


def style_par(par, rstyle, keep_next=False):
    pf = par.paragraph_format
    pf.alignment = ALIGN.get(rstyle.alignment, WD_ALIGN_PARAGRAPH.LEFT)
    pf.space_before = Pt(rstyle.spaceBefore)
    pf.space_after = Pt(rstyle.spaceAfter)
    if rstyle.leading:
        pf.line_spacing = Pt(rstyle.leading)
    if rstyle.leftIndent:
        pf.left_indent = Pt(rstyle.leftIndent)
    if rstyle.firstLineIndent:
        pf.first_line_indent = Pt(rstyle.firstLineIndent)
    if rstyle.backColor is not None:
        set_shading(par._element.get_or_add_pPr(), rstyle.backColor.hexval()[2:].upper())
    if keep_next:
        pf.keep_with_next = True


def add_par(container, p):
    st = p.style
    name = st.name
    if name == "H1":
        par = container.add_paragraph(style="Heading 1")
    elif name == "H2":
        par = container.add_paragraph(style="Heading 2")
    elif p.bulletText:
        par = container.add_paragraph(style="List Bullet")
    else:
        par = container.add_paragraph()
    add_runs(par, p.text, st)
    style_par(par, st, keep_next=name in ("H1", "H2", "IH", "IH2", "ICAP_KEEP"))
    if p.bulletText:
        par.paragraph_format.first_line_indent = None
    return par


def add_image(container, img, caption=None):
    par = container.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after = Pt(2)
    par.paragraph_format.keep_with_next = caption is not None
    par.add_run().add_picture(img.filename, width=Pt(img.drawWidth))
    return par


def add_table(container, tb, ieee):
    data = tb._cellvalues
    widths = tb._colWidths
    nested = any(isinstance(c, (list, tuple)) for row in data for c in row)
    t = container.add_table(rows=len(data), cols=len(data[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    fam = "Times New Roman" if ieee else "Arial"
    size = 7.5 if ieee else 7.8
    if not nested:
        t.style = "Table Grid"
    for ri, row in enumerate(data):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.width = Pt(widths[ci])
            tcPr = cell._element.get_or_add_tcPr()
            if isinstance(val, (list, tuple)):  # celda con figura(s) y pie
                cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)
                for item in val:
                    if isinstance(item, Image):
                        add_image(cell, item, caption=True)
                    elif isinstance(item, Paragraph):
                        add_par(cell, item)
                continue
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(1)
            par.paragraph_format.space_before = Pt(1)
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = par.add_run(html.unescape(str(val)))
            r.font.name = fam
            r._element.rPr.rFonts.set(qn("w:eastAsia"), fam)
            r.font.size = Pt(size)
            if ri == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                set_shading(tcPr, "7B1E3A")
            elif ri % 2 == 0:
                set_shading(tcPr, "F7EEF1")
    # espacio tras la tabla
    gap = container.add_paragraph()
    gap.paragraph_format.space_after = Pt(2)
    gap.paragraph_format.line_spacing = Pt(4)
    return t


def page_field(par):
    for kind, text in (("begin", None), (None, " PAGE "), ("end", None)):
        r = par.add_run()
        r.font.size = Pt(8)
        if kind:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r._element.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; r._element.append(it)


def set_footer(section, label, fam):
    section.footer.is_linked_to_previous = False
    par = section.footer.paragraphs[0]
    for r in list(par.runs):
        r._element.getparent().remove(r._element)
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = par.add_run(f"{label} · pág. ")
    r.font.size = Pt(8); r.font.name = fam; r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    page_field(par)


def page_setup(section, size, margins):
    section.page_width, section.page_height = size
    section.left_margin, section.right_margin, section.top_margin, section.bottom_margin = margins


def set_columns(section, n, space_cm=0.6):
    sectPr = section._sectPr
    for old in sectPr.findall(qn("w:cols")):
        sectPr.remove(old)
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), str(int(space_cm * 567)))
    sectPr.append(cols)


# ------------------------------------------------------------------ sección 1: documento A4
A4 = (Cm(21.0), Cm(29.7)); LETTER = (Cm(21.59), Cm(27.94))
s0 = doc.sections[0]
page_setup(s0, A4, (Cm(2.0), Cm(2.0), Cm(2.2), Cm(2.2)))
set_footer(s0, "Minería de Datos – Práctica 1 · Wine Reviews", "Arial")

for st_name, size, color in (("Heading 1", 15, "7B1E3A"), ("Heading 2", 11.5, "333333")):
    hs = doc.styles[st_name]
    hs.font.name = "Arial"; hs.font.size = Pt(size); hs.font.bold = True
    hs.font.color.rgb = RGBColor.from_string(color)
    hs.element.rPr.rFonts.set(qn("w:ascii"), "Arial"); hs.element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")

ieee_pending, in_ieee = False, False
container = doc


def flatten(items):
    for it in items:
        if isinstance(it, KeepTogether):
            yield from flatten(it._content)
        else:
            yield it


story = list(flatten(br.story))
i = 0
while i < len(story):
    item = story[i]
    nxt = story[i + 1] if i + 1 < len(story) else None
    if isinstance(item, NextPageTemplate):
        ieee_pending = item.action[1] == "ieee"
    elif isinstance(item, PageBreak):
        if ieee_pending:
            sec = doc.add_section(WD_SECTION.NEW_PAGE)
            page_setup(sec, LETTER, (Cm(1.7), Cm(1.7), Cm(1.7), Cm(1.8)))
            set_columns(sec, 1)
            set_footer(sec, "Artículo IEEE short paper", "Times New Roman")
            ieee_pending, in_ieee = False, True
        else:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    elif type(item).__name__ == "_FrameBreak":
        sec = doc.add_section(WD_SECTION.CONTINUOUS)
        set_columns(sec, 2)
    elif isinstance(item, Paragraph):
        add_par(doc, item)
        # una imagen seguida de su pie: dejar el pie junto a la imagen
    elif isinstance(item, Image):
        cap = nxt if isinstance(nxt, Paragraph) and nxt.style.name in ("CAP", "ICAP") else None
        add_image(doc, item, caption=cap)
    elif isinstance(item, Table):
        add_table(doc, item, in_ieee)
    elif isinstance(item, Spacer):
        pass
    i += 1

out = Path(__file__).parent / "Informe_Practica1_Wine_Reviews.docx"
doc.save(out)
print("DOCX:", out)
